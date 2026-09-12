# -*- coding: utf-8 -*-
"""Converte times.xlsx (Guia visual de hunt, por Safnaw) em site/src/data/hunts.json"""
import json, re
import openpyxl

SRC = 'https://docs.google.com/spreadsheets/d/1JcYTCkuKiYK6OcEx9CC-LxRLsPRAPvKaHza-cZuEh1s/edit?gid=68823053#gid=68823053'
wb = openpyxl.load_workbook('times.xlsx', data_only=True)

def clean(v):
    return '' if v is None else str(v).replace('⭐', '').replace('★', '').strip()

def names(v):
    s = clean(v)
    if not s or s in ('—', '-'): return []
    out = []
    for n in re.split(r'\s*,\s*', s):
        n = n.strip()
        if not n: continue
        n = re.sub(r'^S\.\s*', 'Shiny ', n)
        n = n.replace('S.Zard', 'Shiny Charizard').replace('Shiny Zard', 'Shiny Charizard')
        n = n.replace('SevIper', 'Seviper').replace('Eletric', 'Electric')
        out.append(n)
    return out

def parse_cards(sheet):
    ws = wb[sheet]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    cards = []
    for i, r in enumerate(rows):
        for c in (0, 3, 6):
            v = clean(r[c]) if c < len(r) else ''
            if v.startswith('★') or (r[c] and str(r[c]).startswith('★')):
                title = v.lstrip('★ ').strip().title()
                g = lambda dr, dc: rows[i + dr][c + dc] if i + dr < len(rows) and c + dc < len(rows[i + dr]) else None
                cards.append({
                    'name': title,
                    'hunt': clean(g(1, 0)).replace('Hunt:', '').strip(),
                    'tanks': names(g(3, 0)), 'dps': names(g(3, 1)),
                    'otherTanks': names(g(5, 0)), 'otherDps': names(g(5, 1)),
                    'smeargle': [s.strip() for s in clean(g(6, 0)).replace('🎨', '').replace('Smeargle:', '').split(',') if s.strip()],
                })
    return cards

elements = parse_cards('Cards - Elementos')
hunts = parse_cards('Cards - Hunts')
ws = wb['Cards - Elementos']
notes = [clean(ws.cell(1, 9).value), clean(ws.cell(2, 9).value)]

out = {'source': SRC, 'author': 'Safnaw', 'twitch': 'https://www.twitch.tv/georgezanelato', 'notes': [n for n in notes if n],
       'elements': elements, 'hunts': hunts}
json.dump(out, open('site/src/data/hunts.json', 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
print('elements', [e['name'] for e in elements])
print('hunts', [h['name'] for h in hunts])
print(elements[0])
